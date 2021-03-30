import docx
import csv

def to_doc(in_path):
	raw_file = open(in_path)
	csv_file =  list(csv.reader(raw_file))
	doc = docx.Document()
	n_cols=max([len(csv_i) for csv_i in csv_file])
	n_rows=len(csv_file)+1
	table = doc.add_table(rows=n_rows,cols=n_cols)
	for i,line_i in enumerate(csv_file):
#		print(len(line_i))
		for j,line_j in enumerate(line_i):
#			print(j)
			cell_i=table.cell(i+1,j)
			cell_i.text=line_j
	doc.save('helloworld.docx')

to_doc("MSR.csv")